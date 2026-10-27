#!/usr/bin/env python3
"""Derive auditable tables, scientific figures and an HTC comparison Word report.

Use --analyze with plotting Python; use --docx with the bundled document Python.
"""
import argparse
from collections import defaultdict
import csv
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import re
import statistics

REPO=Path(__file__).resolve().parents[2]
CASE_ORDER=['free','following','signal','initial','progressive']
STYLE={'sumo':('SUMO','#333333'),'discrete-event-1':('Conservative 1 JVM','#247ba0'),
       'time-warp-1':('Optimistic 1 JVM','#ed9b40'),'discrete-event-2':('Conservative 2 JVMs','#006d77'),
       'time-warp-2':('Optimistic 2 JVMs','#bc4b51'),'reference':('Fixed placement','#247ba0'),
       'migration':('Rebalancing enabled','#ed9b40')}


def read(p):return json.loads(Path(p).read_text())
def write(p,v):Path(p).write_text(json.dumps(v,indent=2,allow_nan=False)+'\n')
def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def median(v):return statistics.median(v) if v else None

def describe(values):
    return dict(n=len(values),median=median(values),minimum=min(values) if values else None,maximum=max(values) if values else None)


def stamp(line):
    m=re.search(r'(\d{4}-\d\d-\d\d \d\d:\d\d:\d\d[,.]\d+)',line)
    return datetime.fromisoformat(m[1].replace(',','.')) if m else None


def intervals(log):
    starts=[];ends=[]
    for line in log.splitlines():
        at=stamp(line)
        if not at:continue
        if 'simulation_start_request_sent_to_gtm' in line:starts.append(at)
        if '[CONSERVATIVE-COMPLETE]' in line or '[ACCOUNTED-COMPLETE]' in line:ends.append(at)
    return (max(ends)-min(starts)).total_seconds() if starts and ends else None


def derive(study,out):
    out.mkdir(parents=True,exist_ok=True)
    data=read(study/'results.json');manifest=read(study/'manifest.json')
    assert data['manifestSha256']==digest(study/'manifest.json')
    assert data['allAttemptsAccountedFor']
    rss_audit=read(study/'sumo-rss-audit/results.json')
    assert rss_audit['allTripRecordsExact'] and rss_audit['allAttemptsAccountedFor']
    flat=[]
    for record in data['records']:
        r=record['result'];m=r.get('measurement',{});rt=m.get('runtime',{})
        log=Path(record['directory'])/'htc/runtime.log'
        body=log.read_text() if log.exists() else ''
        interval=intervals(body) if body else None
        cpus=[float(v) for v in m.get('cpuSecondsByPid',{}).values()]
        from benchmark_metrics import imbalance
        loads=[v for v in rt.get('loads',[]) if v['imbalance']]
        row=dict(config=record['config'],case=record['case'],kind=record['kind'],repeat=record['repeat'],
            warmup=record['warmup'],accepted=bool(r.get('passed')),engineCompleted=r.get('engineCompleted'),
            attemptWallSeconds=record['attemptWallSeconds'],
            wallSeconds=m.get('completionSeconds'),cpuSeconds=m.get('cpuSeconds'),
            peakRssMiB=m.get('peakAggregateRssBytes',0)/1024**2,
            requestToCompletionSeconds=interval,rollbacks=rt.get('rollbacks',0),
            verifiedShardMoves=len(rt.get('verifiedMoves',[])),migrationWaves=len(rt.get('waves',[])),
            pauseWindows=len(rt.get('pauses',[])),pauseSeconds=sum(p['totalSeconds'] for p in rt.get('pauses',[])),
            restoredEntities=sum(v['entities'] for v in rt.get('waves',[])),
            pauseRequests=body.count('Requesting TimeManager to pause for migration of shards:'),
            freezeDispatches=body.count('[MIGRATION-FREEZE]'),
            declinedRequests=len(re.findall(r"Migration request '[^']+' declined:",body)),
            liveLoadSnapshots=len(loads),firstLoadCv=loads[0]['imbalance']['cv'] if loads else None,
            lastLoadCv=loads[-1]['imbalance']['cv'] if loads else None,
            cpuWorkCv=(imbalance(cpus) or {}).get('cv') if cpus else None,
            activeCarRestores=len(r.get('physicalChecks',{}).get('activeCarRestores',[])),
            parity=r.get('reportParity'),failure=r.get('exception') or r.get('physicalChecks',{}).get('failure'),
            source=record['directory'])
        if record['kind']=='sumo':
            supplement=next(v for v in rss_audit['records'] if v['case']==record['case'] and v['repeat']==record['repeat'])
            row['originalForkHighWaterRssMiB']=row['peakRssMiB']
            row['peakRssMiB']=supplement['resource']['peakRssKiB']/1024
            row['memorySource']='separately retained GNU-time RSS audit; original elapsed and CPU unchanged'
        if record['kind']=='traffic':
            p=Path(record['directory'])/'physical-comparison.json'
            if p.exists():
                d=read(p);vs=d['vehicles']
                row.update(tripTimeMaeSeconds=statistics.mean(abs(v['durationDelta']) for v in vs),
                    maxAbsTripTimeSeconds=max(abs(v['durationDelta']) for v in vs),
                    routeDistanceRmseM=statistics.mean(v['routeDistanceRmseM'] for v in vs),
                    speedRmseMps=statistics.mean(v['speedRmseMps'] for v in vs),vehicles=len(vs))
        flat.append(row)
    keys=list(dict.fromkeys(k for row in flat for k in row))
    with (out/'all-attempts.csv').open('w') as stream:
        w=csv.DictWriter(stream,fieldnames=keys);w.writeheader();w.writerows(flat)
    groups=[]
    measured=[r for r in flat if not r['warmup']]
    for cfg in manifest['configurations']:
        rows=[r for r in measured if r['config']==cfg['id']]
        ok=[r for r in rows if r['accepted']]
        groups.append(dict(**cfg,attempts=len(rows),accepted=len(ok),
            metrics={k:describe([r[k] for r in ok if r.get(k) is not None]) for k in (
                'wallSeconds','cpuSeconds','peakRssMiB','requestToCompletionSeconds','rollbacks','verifiedShardMoves',
                'pauseSeconds','pauseRequests','freezeDispatches','declinedRequests','migrationWaves','restoredEntities','firstLoadCv','lastLoadCv','cpuWorkCv','activeCarRestores')}))
    ratios=[]
    for case in CASE_ORDER:
        comparisons=[('time-warp-1','discrete-event-1'),('time-warp-2','discrete-event-2'),
                     ('discrete-event-2','discrete-event-1'),('time-warp-2','time-warp-1'),
                     ('discrete-event-1','sumo'),('time-warp-1','sumo')] if case in CASE_ORDER[:3] else [('migration','reference')]
        for numerator,denominator in comparisons:
            pairs=[]
            for repeat in range(1,6):
                a=next((r for r in measured if r['repeat']==repeat and r['config']==case+'-'+numerator and r['accepted']),None)
                b=next((r for r in measured if r['repeat']==repeat and r['config']==case+'-'+denominator and r['accepted']),None)
                if a and b and a['wallSeconds'] is not None and b['wallSeconds']:
                    pairs.append(dict(repeat=repeat,ratio=a['wallSeconds']/b['wallSeconds']))
            ratios.append(dict(case=case,numerator=numerator,denominator=denominator,pairs=pairs,
                               summary=describe([p['ratio'] for p in pairs])))
    summary=dict(campaignAttemptSeconds=sum(r['attemptWallSeconds'] for r in flat),
        warmupAttemptSeconds=sum(r['attemptWallSeconds'] for r in flat if r['warmup']),
        measuredAttemptSeconds=sum(r['attemptWallSeconds'] for r in flat if not r['warmup']),
        date='2026-10-27',manifestSha256=digest(study/'manifest.json'),
        attempts=len(flat),measuredAttempts=len(measured),warmupAttempts=len(flat)-len(measured),
        acceptedMeasured=sum(r['accepted'] for r in measured),
        failedMeasured=[r for r in measured if not r['accepted']],groups=groups,ratios=ratios,
        allAttemptsAccountedFor=data['allAttemptsAccountedFor'],
        rssAudit=dict(attempts=len(rss_audit['records']),measuredAttempts=sum(not v['warmup'] for v in rss_audit['records']),
            allTripRecordsExact=rss_audit['allTripRecordsExact'],manifestSha256=rss_audit['manifestSha256']),
        measuredParityChecks=[p for p in data['parity'] if p['repeat']>0],
        sourceRecords={str(p.relative_to(REPO)):digest(p) for p in [
            REPO/'docs/validation/micro-discharge-2026-10-25.json',
            REPO/'docs/validation/micro-scientific-2026-10-25.json',
            REPO/'docs/validation/micro-field-v2-2026-10-26.json',
            REPO/'docs/validation/load-balancing-completion-2026-10-26.json',
            REPO/'docs/validation/micro-spillback-2026-10-26.json']})
    write(out/'summary.json',summary)
    write(out/'all-attempts.json',flat)
    with (out/'configuration-summary.csv').open('w') as stream:
        fields=['config','attempts','accepted','wallMedianSeconds','wallMinSeconds','wallMaxSeconds','cpuMedianSeconds','peakRssMedianMiB','movesMedian','pauseMedianSeconds']
        w=csv.DictWriter(stream,fieldnames=fields);w.writeheader()
        for g in groups:
            m=g['metrics'];w.writerow(dict(config=g['id'],attempts=g['attempts'],accepted=g['accepted'],
                wallMedianSeconds=m['wallSeconds']['median'],wallMinSeconds=m['wallSeconds']['minimum'],wallMaxSeconds=m['wallSeconds']['maximum'],
                cpuMedianSeconds=m['cpuSeconds']['median'],peakRssMedianMiB=m['peakRssMiB']['median'],movesMedian=m['verifiedShardMoves']['median'],pauseMedianSeconds=m['pauseSeconds']['median']))
    return summary,flat,manifest,data


def figures(study,out,summary,flat,manifest,data):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import numpy as np
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,
        'axes.spines.right':False,'figure.facecolor':'white','savefig.facecolor':'white'})
    folder=out/'figures';folder.mkdir(exist_ok=True)
    saved=[]
    def save(fig,name):
        fig.savefig(folder/(name+'.png'),dpi=190,bbox_inches='tight')
        fig.savefig(folder/(name+'.svg'),bbox_inches='tight');plt.close(fig);saved.append(name)
    def legend_key(cfg):return cfg.split('-',1)[1]
    def points(ax,configs,metric,log=False):
        for x,cfg in enumerate(configs):
            rows=[r for r in flat if r['config']==cfg and not r['warmup'] and r['accepted'] and r.get(metric) is not None]
            vals=[r[metric] for r in rows]
            color=STYLE[legend_key(cfg)][1]
            if vals:
                ax.scatter([x+(r['repeat']-3)*.045 for r in rows],vals,c=color,s=20,alpha=.8)
                ax.plot([x-.25,x+.25],[median(vals)]*2,color=color,lw=2.5)
            else:ax.text(x,.03,'No accepted runs',rotation=90,ha='center',va='bottom',transform=ax.get_xaxis_transform(),fontsize=8)
        short={'sumo':'SUMO','discrete-event-1':'C1','time-warp-1':'O1','discrete-event-2':'C2','time-warp-2':'O2','reference':'Fixed','migration':'Rebalance'}
        ax.set_xticks(range(len(configs)),[short[legend_key(c)] for c in configs],fontsize=9)
        if log:ax.set_yscale('log')
        ax.grid(axis='y',alpha=.2)
    fig,axes=plt.subplots(1,3,figsize=(11,3.9))
    for ax,case in zip(axes,CASE_ORDER[:3]):
        cfgs=[case+'-'+k for k in list(STYLE)[:5]]
        points(ax,cfgs,'wallSeconds',True);ax.set_title(case.capitalize());ax.set_ylabel('Launch to completion (s, log scale)')
    fig.suptitle('Output inclusive execution time   Five measured attempts per configuration',y=1.04)
    fig.tight_layout();save(fig,'execution-times')
    fig,axes=plt.subplots(2,3,figsize=(11,6.5))
    for col,case in enumerate(CASE_ORDER[:3]):
        cfgs=[case+'-'+k for k in list(STYLE)[:5]]
        for row,metric,label in [(0,'cpuSeconds','Process CPU (s)'),(1,'peakRssMiB','Peak process RSS (MiB)')]:
            points(axes[row,col],cfgs,metric,True);axes[row,col].set_title(case.capitalize());axes[row,col].set_ylabel(label+'   log scale')
    fig.tight_layout();save(fig,'resources')
    fig,axes=plt.subplots(1,3,figsize=(11,3.4))
    for ax,case in zip(axes,CASE_ORDER[:3]):
        cfgs=[case+'-'+k for k in list(STYLE)[1:5]]
        points(ax,cfgs,'rollbacks');ax.set_title(case.capitalize());ax.set_ylabel('Rollback log entries per run')
    fig.tight_layout();save(fig,'rollbacks')
    fig,axes=plt.subplots(1,2,figsize=(10,3.6))
    for ax,case in zip(axes,CASE_ORDER[3:]):
        points(ax,[case+'-reference',case+'-migration'],'wallSeconds');ax.set_title(case.capitalize());ax.set_ylabel('Launch to completion (s)')
    fig.tight_layout();save(fig,'balancing-times')
    fig,axes=plt.subplots(1,3,figsize=(11,3.4))
    for ax,metric,label in zip(axes,['verifiedShardMoves','pauseSeconds','activeCarRestores'],
                             ['Verified shard moves','Global pause duration (s)','Active Car restore observations']):
        points(ax,[c+'-'+p for c in CASE_ORDER[3:] for p in ('reference','migration')],metric)
        ax.set_xticklabels(['Initial\nfixed','Initial\nrebalance','Progressive\nfixed','Progressive\nrebalance'],fontsize=8)
        ax.set_ylabel(label)
    fig.tight_layout();save(fig,'rebalancing-costs')
    fig,axes=plt.subplots(2,2,figsize=(10,6))
    for col,case in enumerate(CASE_ORDER[3:]):
        for placement in ('reference','migration'):
            record=next(r for r in data['records'] if r['repeat']==1 and r['config']==case+'-'+placement)
            samples=record['result'].get('measurement',{}).get('runtime',{}).get('loads',[])
            timed=[(stamp(v['timestamp']),v) for v in samples if v.get('timestamp')]
            if timed:
                origin=timed[0][0];ts=[(t-origin).total_seconds() for t,v in timed]
                color=STYLE[placement][1]
                for node in range(len(timed[0][1]['values'])):
                    axes[0,col].step(ts,[v['values'][node] for t,v in timed],where='post',color=color,ls='-' if node==0 else '--',label=f'{placement} node {node}')
                axes[1,col].step(ts,[v['imbalance']['cv'] if v['imbalance'] else math.nan for t,v in timed],where='post',color=color,label=STYLE[placement][0])
        axes[0,col].set_title(case.capitalize()+'   First measured repetition')
        axes[0,col].set_ylabel('Live entities per node');axes[0,col].legend(fontsize=7)
        axes[1,col].set_ylabel('Live entity load CV');axes[1,col].set_xlabel('Seconds since first observation in each run')
        axes[1,col].legend(fontsize=8)
    fig.tight_layout();save(fig,'live-loads')
    fig,axes=plt.subplots(1,2,figsize=(10,3.8))
    for ax,case in zip(axes,CASE_ORDER[3:]):
        record=next(r for r in data['records'] if r['repeat']==1 and r['config']==case+'-migration')
        ps=record['result'].get('measurement',{}).get('runtime',{}).get('pauses',[])
        if ps:
            y=np.arange(len(ps));left=np.zeros(len(ps))
            for key,label,color in [('freezeSeconds','Owner freeze','#247ba0'),('relocationSeconds','Relocation and window ACKs','#ed9b40'),('resumeSeconds','Owner resume','#006d77')]:
                vals=[p[key] for p in ps];ax.barh(y,vals,left=left,label=label,color=color);left+=vals
            ax.set_yticks(y,[f'Pause {i+1}' for i in y]);ax.legend(fontsize=7)
        else:ax.text(.5,.5,'No completed pause window',ha='center',transform=ax.transAxes)
        ax.set_title(case.capitalize()+'   First measured repetition');ax.set_xlabel('Elapsed pause time (s)')
    fig.tight_layout();save(fig,'migration-phases')
    # All elapsed comparisons retain individual repeat pairing.
    selected=[r for r in summary['ratios'] if r['denominator']!='sumo']
    fig,ax=plt.subplots(figsize=(10,5))
    for y,r in enumerate(selected):
        vals=[p['ratio'] for p in r['pairs']]
        if vals:
            ax.plot([min(vals),max(vals)],[y,y],color='#777777');ax.scatter(vals,[y]*len(vals),s=16,color='#247ba0');ax.scatter([median(vals)],[y],marker='D',s=28,color='#bc4b51')
    ax.axvline(1,color='black',ls='--',lw=1)
    ax.set_yticks(range(len(selected)),[f"{r['case']}: {r['numerator']} / {r['denominator']}" for r in selected],fontsize=8)
    ax.set_xlabel('Paired elapsed time ratio   Above 1 means numerator took longer')
    ax.invert_yaxis();ax.grid(axis='x',alpha=.2);fig.tight_layout();save(fig,'paired-ratios')
    # Physical results of the predeclared first measured conservative run.
    comparisons={}
    for case in CASE_ORDER[:3]:
        run=study/'runs'/f'r1-{case}-discrete-event-1'
        p=run/'physical-comparison.json'
        if p.exists():comparisons[case]=read(p)
    fig,axes=plt.subplots(1,3,figsize=(11,3.6))
    for ax,case in zip(axes,CASE_ORDER[:3]):
        vs=comparisons.get(case,{}).get('vehicles',[])
        if vs:
            a=[v['sumoDuration'] for v in vs];b=[v['htcDuration'] for v in vs]
            ax.scatter(a,b,color='#247ba0');lo=min(a+b);hi=max(a+b)
            if hi-lo<.1:lo-=1;hi+=1
            ax.plot([lo,hi],[lo,hi],color='black',ls='--');ax.set_xlim(lo-.2,hi+.2);ax.set_ylim(lo-.2,hi+.2)
        ax.set_title(case.capitalize());ax.set_xlabel('SUMO travel time (s)');ax.set_ylabel('HTC travel time (s)')
    fig.tight_layout();save(fig,'travel-times')
    from validate_sumo_micro import trajectory, sumo_samples
    fig,axes=plt.subplots(2,2,figsize=(10,6))
    for col,case in enumerate(('following','signal')):
        root=study/'runs'/f'r1-{case}-discrete-event-1'
        rows=[json.loads(line) for p in (root/'htc/reports').rglob('*.jsonl') for line in p.read_text().splitlines() if line]
        canon=read(study/'inputs'/case/'canonical.json')
        native=trajectory(rows,canon);reference=sumo_samples(study/'runs'/f'r1-{case}-sumo/sumo/fcd.xml',canon)
        vehicle=canon['vehicles'][-1]
        for values,label,color in [(native,'HTC conservative','#247ba0'),(reference,'SUMO','#333333')]:
            series=sorted((t,v) for (vid,t),v in values.items() if vid==vehicle)
            for row,key in [(0,'distance'),(1,'speed')]:
                axes[row,col].plot([t for t,v in series],[v[key] for t,v in series],label=label,color=color,lw=1.4,ls='-' if label.startswith('HTC') else '--')
        axes[0,col].set_title(case.capitalize()+f'   Last scheduled vehicle {vehicle}')
        axes[0,col].set_ylabel('Route distance (m)');axes[1,col].set_ylabel('Speed (m/s)');axes[1,col].set_xlabel('Simulation time (s)');axes[0,col].legend(fontsize=8)
    fig.tight_layout();save(fig,'trajectories')
    # Queue and matched discharge statistics already use the frozen comparator definitions.
    d=comparisons.get('signal',{})
    write(out/'representative-physical.json',comparisons)
    # Report exact measured intervals and label the residual rather than inventing a simulation-only phase.
    fig,ax=plt.subplots(figsize=(10,4))
    cfgs=[g for g in summary['groups'] if g['kind']=='traffic']
    labels=[]
    for y,g in enumerate(cfgs):
        total=g['metrics']['wallSeconds']['median'];interval=g['metrics']['requestToCompletionSeconds']['median']
        if total is not None and interval is not None:
            ax.barh(y,max(0,total-interval),color='#aaaaaa',label='Launch and remaining notification time' if y==0 else None)
            ax.barh(y,interval,left=max(0,total-interval),color='#247ba0',label='Start request through runtime completion' if y==0 else None)
        labels.append(g['id'])
    ax.set_yticks(range(len(cfgs)),labels,fontsize=8);ax.invert_yaxis();ax.set_xlabel('Median elapsed seconds   Interval medians are descriptive')
    ax.legend(fontsize=8);fig.tight_layout();save(fig,'execution-intervals')
    from sumo_scientific_metrics import traffic_metrics
    scientific={}
    for case,d in comparisons.items():
        view=out/'analysis-views'/case
        view.mkdir(parents=True,exist_ok=True)
        for name,target in [('htc',study/'runs'/f'r1-{case}-discrete-event-1/htc'),('sumo',study/'runs'/f'r1-{case}-sumo/sumo')]:
            link=view/name
            if not link.exists():link.symlink_to(target,target_is_directory=True)
        scientific[case]=traffic_metrics(view,dict(d,canonical=read(study/'inputs'/case/'canonical.json')),
                                        read(REPO/'scripts/scenario/sumo_scientific_protocol_v2.json'))
    write(out/'representative-scientific.json',scientific)
    case='signal';root=study/'runs'/f'r1-{case}-discrete-event-1'
    rows=[json.loads(line) for p in (root/'htc/reports').rglob('*.jsonl') for line in p.read_text().splitlines() if line]
    canon=read(study/'inputs'/case/'canonical.json')
    a=trajectory(rows,canon);b=sumo_samples(study/'runs'/f'r1-{case}-sumo/sumo/fcd.xml',canon)
    fig,axes=plt.subplots(1,2,figsize=(10,3.7))
    from collections import Counter
    times=[round(i*.1,9) for i in range(1801)]
    for samples,label,color in [(a,'HTC','#247ba0'),(b,'SUMO','#333333')]:
        counts=Counter(t for (vid,t),v in samples.items() if v['link']==canon['links'][0] and v['speed']<.1)
        axes[0].step(times,[counts[t] for t in times],where='post',label=label,color=color,lw=1.2,ls='-' if label=='HTC' else '--')
    pairs=scientific.get('signal',{}).get('dischargeHeadways',{}).get('pairs',[])
    for label,key,color,offset in [('HTC','htcSeconds','#247ba0',-.18),('SUMO','sumoSeconds','#333333',.18)]:
        axes[1].bar(np.arange(len(pairs))+offset,[p[key] for p in pairs],width=.34,label=label,color=color)
    axes[0].set_xlabel('Simulation time (s)');axes[0].set_ylabel('Approach vehicles below 0.1 m/s');axes[0].legend(fontsize=8)
    axes[1].set_xlabel('Matched consecutive vehicle pair within green');axes[1].set_ylabel('Discharge headway (s)');axes[1].legend(fontsize=8)
    fig.tight_layout();save(fig,'queues-headways')
    fig,axes=plt.subplots(1,2,figsize=(10,3.7))
    for ax,case in zip(axes,CASE_ORDER[3:]):
        for placement in ('reference','migration'):
            samples=read(study/'runs'/f'r1-{case}-{placement}/htc/process-samples.json')
            times=[];cvs=[]
            from benchmark_metrics import imbalance
            for a,b in zip(samples,samples[1:]):
                if len(a['nodes'])!=2 or len(b['nodes'])!=2 or not all(a['nodes']) or not all(b['nodes']):continue
                if [n['pid'] for n in a['nodes']] != [n['pid'] for n in b['nodes']]:continue
                work=[max(0,y['cpuSeconds']-x['cpuSeconds']) for x,y in zip(a['nodes'],b['nodes'])]
                metric=imbalance(work)
                times.append(b['seconds']);cvs.append(metric['cv'] if metric else math.nan)
            ax.plot(times,cvs,color=STYLE[placement][1],label=STYLE[placement][0],alpha=.8,lw=.9)
        ax.set_title(case.capitalize()+'   First measured repetition');ax.set_xlabel('Seconds since launch');ax.set_ylabel('Per interval CPU work CV');ax.legend(fontsize=8)
    fig.tight_layout();save(fig,'cpu-imbalance')
    return saved


def document(study,out):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    summary=read(out/'summary.json');flat=read(out/'all-attempts.json');manifest=read(study/'manifest.json')
    physical=read(out/'representative-physical.json');scientific=read(out/'representative-scientific.json')
    field=read(REPO/'docs/validation/micro-field-v2-2026-10-26.json')
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
    for name,size in [('Title',24),('Subtitle',13),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_after=Pt(10);s.paragraph_format.space_before=Pt(10)
    doc.styles['Caption'].font.name='Times New Roman';doc.styles['Caption'].font.size=Pt(9)
    doc.styles['Caption'].font.color.rgb=RGBColor(50,50,50)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('HTC comparison study  |  ');r.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    def p(text,bold=False):
        para=doc.add_paragraph();run=para.add_run(text);run.bold=bold;return para
    def page(title):
        h=doc.add_heading(title,level=1);h.paragraph_format.page_break_before=True
    def heading(title):doc.add_heading(title,level=2)
    def table(headers,rows,widths=None):
        t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        if widths:
            for col,width in zip(t.columns,widths):col.width=Inches(width)
        for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h)
        for row in rows:
            cells=t.add_row().cells
            for i,value in enumerate(row):cells[i].text=str(value)
        for index,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr();keep=OxmlElement('w:cantSplit');pr.append(keep)
            if index==0:
                repeat=OxmlElement('w:tblHeader');pr.append(repeat)
            for col,cell in enumerate(row.cells):
                if widths:cell.width=Inches(widths[col])
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr()
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'263746' if index==0 else ('F1F4F6' if index%2==0 else 'FFFFFF'));tcpr.append(shade)
                borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                tcpr.append(borders)
                margins=OxmlElement('w:tcMar')
                for edge in ('top','left','bottom','right'):
                    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');margins.append(e)
                tcpr.append(margins)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.02
                    if col>0 and (index==0 or len(cell.text)<65):para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.size=Pt(9);run.font.bold=index==0
                        run.font.color.rgb=RGBColor(255,255,255) if index==0 else RGBColor(0,0,0)
        spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=Pt(2);spacer.add_run().font.size=Pt(2)
        return t
    figure_number=0
    def fig(name,caption,width=6.9):
        nonlocal figure_number
        figure_number+=1
        path=out/'figures'/(name+'.png')
        para=doc.add_paragraph();para.paragraph_format.space_after=Pt(3);para.paragraph_format.keep_with_next=True
        shape=para.add_run().add_picture(str(path),width=Inches(width))
        shape._inline.docPr.set('descr',caption)
        c=doc.add_paragraph(f'Figure {figure_number}. {caption}',style='Caption')
        c.paragraph_format.space_after=Pt(9)
    def number(value,dec=2):return 'Unavailable' if value is None else f'{value:,.{dec}f}'
    def span(d):
        return 'Unavailable' if not d['n'] else f"{number(d['median'])} [{number(d['minimum'])}, {number(d['maximum'])}]"
    def group(config):return next(g for g in summary['groups'] if g['id']==config)
    def ratio(case,a,b):return next(r for r in summary['ratios'] if r['case']==case and r['numerator']==a and r['denominator']==b)
    accepted=summary['acceptedMeasured'];total=summary['measuredAttempts']
    title=doc.add_paragraph('HTC traffic fidelity and runtime performance',style='Title')
    doc.add_paragraph('Comparisons with SUMO and between execution and balancing policies',style='Subtitle')
    p('Research evidence report  |  27 October 2026')
    heading('Executive findings')
    p(f'The bounded campaign accounts for {summary["attempts"]} attempts: {summary["warmupAttempts"]} excluded warm-ups and {total} measured attempts. {accepted} measured attempts satisfy their runtime and physical acceptance checks. The results describe small workloads on one host with real TCP between JVMs. They do not establish city-scale capacity or performance on separate machines.')
    p('HTC and SUMO are compared on matched longitudinal traffic cases. Conservative and optimistic HTC runs must produce the same committed physical output. Rebalancing is assessed independently under conservative execution because optimistic migration is not supported. The report distinguishes traffic agreement from execution cost and from empirical validity.')
    rows=[]
    for case in CASE_ORDER[:3]:
        g=group(case+'-discrete-event-1');o=group(case+'-time-warp-1');s=group(case+'-sumo')
        rows.append([case.capitalize(),number(s['metrics']['wallSeconds']['median']),number(g['metrics']['wallSeconds']['median']),number(o['metrics']['wallSeconds']['median'])])
    table(['Case','SUMO median s','Conservative 1 JVM s','Optimistic 1 JVM s'],rows,[1.4,1.5,2,2])
    p('Medians summarize five attempts; the detailed results retain individual values and acceptance failures. Small SUMO CLI cases and distributed HTC have different startup and output costs, so end-to-end ratios describe these configured executions rather than the relative efficiency of their car-following formulas.')
    p(f'Total recorded campaign attempt time is {summary["campaignAttemptSeconds"]/60:.1f} minutes, including {summary["warmupAttemptSeconds"]/60:.1f} minutes of excluded warm-ups and harness cleanup. Measured-attempt time is {summary["measuredAttemptSeconds"]/60:.1f} minutes. These campaign costs differ from the pre-cleanup execution values used in performance ratios.')
    p('The field study remains inconclusive. Its eight following episodes and four queues pass their comparative checks, but fall short of the frozen targets of ten and five. SUMO agreement and repeatable engine output do not close that empirical evidence gap.')

    page('Study design and experimental controls')
    p('The benchmark was specified before measured execution. The immutable manifest records 19 configurations, five measured repetitions, one excluded warm-up per configuration, randomized order seed 20261027, a 300-second per-run limit, executable and source hashes, generated inputs, commands and host details. Every attempt is retained. Warm-ups start fresh JVMs; they warm operating-system caches and exercise the fixture, rather than preserving a compiled JVM between measured runs.')
    table(['Workload','Configuration matrix','Measured attempts'],[
        ['Free flow','SUMO plus C1 O1 C2 O2',25],['Following','SUMO plus C1 O1 C2 O2',25],
        ['Signal right turn','SUMO plus C1 O1 C2 O2',25],['Initial balancing','Fixed placement and rebalancing on 2 JVMs',10],
        ['Progressive balancing','Fixed placement and rebalancing on 2 JVMs',10]], [1.5,3.8,1.6])
    p('C and O denote conservative and optimistic execution; the suffix is the JVM count. The free case has one Car, the following case four Cars, and the signal case eighteen Cars. All use a 0.1-second physics/runtime clock, static routes, deterministic drivers and the existing default behavioral coefficients. The signal fixture is the released v2 right-turn case, not a fresh scientific holdout.')
    cpu=manifest['host']['cpu'];model=re.search(r'model name\s*:\s*(.*)',cpu)
    p(f'The host CPU is {model[1] if model else "recorded in the manifest"}. Four logical CPUs {manifest["cpuIds"]} are assigned in total: four to one JVM or two disjoint pairs to two JVMs. Total maximum Java heap is 4 GiB, split equally for two JVMs; initial heap remains 128 MiB per JVM. SUMO receives the same allowed CPU set but does not thereby become a four-thread execution. Java active processor count follows the assigned CPUs.')
    p('Configurations run sequentially. Other host applications are not terminated; load averages and system CPU counters are retained to expose contention. Affinity limits the study processes but does not reserve exclusive cores. Both JVMs share the same host, filesystem and loopback TCP. SUMO 1.27.1, Java version, operating system and full commands are recorded in the evidence manifest.')
    raw_records=read(study/'results.json')['records']
    observations=[r['result'].get('measurement',{}) for r in raw_records if not r['warmup']]
    host_low=min(v['hostLoad1Min'] for v in observations if v.get('hostLoad1Min') is not None)
    host_high=max(v['hostLoad1Max'] for v in observations if v.get('hostLoad1Max') is not None)
    p(f'The retained one-minute host load averages span {host_low:.2f} to {host_high:.2f} during measured runs. Load includes runnable and uninterruptible tasks and is not a process CPU-utilization percentage. These high shared-host loads limit attribution of elapsed outliers to the simulator or policy alone.')
    p('HTC writes committed JSON reports, including every MICRO update, through its reporter. SUMO writes FCD, tripinfo and summary XML. Output remains enabled for verification; formats, event volumes and flushing are not identical. Timing therefore includes each engine’s configured reporting pipeline.')

    page('Metrics and interpretation')
    table(['Metric','Definition and unit'],[
        ['Travel duration','Simulated arrival minus actual departure, in physical seconds.'],
        ['Execution time','Monotonic launch to acknowledged completion and stop notification; harness teardown excluded. SUMO ends at process completion.'],
        ['Process CPU','User plus system CPU-seconds. HTC uses final live process counters; SUMO uses terminal accounting for CPU and a separate GNU-time audit for native RSS.'],
        ['Peak memory','Maximum simultaneous sampled RSS sum for HTC; native process high-water RSS from the GNU-time audit for SUMO, in MiB.'],
        ['Load imbalance','Population standard deviation divided by mean per-node load. Also retain maximum divided by mean; all-zero loads are unavailable.'],
        ['Pause duration','Monotonic interval from owner-freeze dispatch through the last matching resume acknowledgement.'],
        ['Rollback count','Number of TW-ROLLBACK log entries, including work later replayed; not a count of incorrect committed trips.'],
        ['Paired time ratio','Numerator elapsed time divided by denominator for the same workload and repetition block. Above one means longer execution.']], [1.6,5.3])
    p('Process samples and available Prometheus metrics are collected externally every 250 ms. RSS is sampled, so short peaks can be missed. Java RSS includes heap, stacks and native memory; the configured heap limit is not an RSS limit. Summing per-process peak RSS would invent a simultaneous peak and is deliberately avoided. The short SUMO executions use kernel terminal accounting because quarter-second sampling alone would under-observe them. A separate 18-attempt RSS audit avoids the large Python fork image inherited by the original launcher. It repeats all three SUMO cases once for warm-up and five times for measurement, preserves every original value, and reproduces all original trip records exactly. Only the native RSS series comes from that supplement; original elapsed and CPU measurements are unchanged.')
    p('Median and minimum-to-maximum ranges describe execution variability; points show each accepted measured run. Repetitions share deterministic traffic and one host, so they estimate execution variability rather than independent traffic realizations. No significance or city-scale speedup is inferred. Failed attempts stay in the ledger and do not enter accepted-run medians.')
    p('Load counters represent live entities and can change as cohorts enroll or retire. Entity balance is not CPU balance. A lower entity-count CV alone does not demonstrate faster execution; pause overhead, transport and speculative work must also be considered.')

    page('Traffic agreement with SUMO')
    fig('travel-times','Trip durations for every vehicle in the first measured conservative one-JVM run. The dashed line is exact agreement; all comparisons use actual departure and arrival records.')
    rows=[]
    for case in CASE_ORDER[:3]:
        s=scientific.get(case,{})
        rows.append([case.capitalize(),number(s.get('meanAbsoluteDurationErrorSeconds'),3),number(s.get('routeDistanceRmseM'),3),number(s.get('speedRmseMps'),3),number(100*s.get('completionCoverage',0),0)+'%'])
    table(['Case','Trip time MAE s','Route RMSE m','Speed RMSE m/s','Coverage'],rows,[1.3,1.4,1.4,1.4,1.4])
    p('Trajectory errors use matched physical timestamps and route coordinates. Comparing local positions from unrelated links would produce a misleading error at transitions, so route-distance coordinates carry cumulative road and connector lengths. No fitted time shift is applied. Trip duration uses explicit terminal records rather than the overlap of observed samples. [1,2]')
    p('Physical agreement is conditional on matched deterministic model inputs and these geometries. The new timing repetitions reuse released cases and cannot supply another independent scientific acceptance claim. Full per-run physical metrics and committed-output comparisons accompany the performance tables.')

    page('Vehicle trajectories and speed')
    fig('trajectories','Route distance and speed of the last scheduled vehicle in the following and signal cases, first measured conservative one-JVM run. Vehicle selection was fixed by schedule order, not by prediction error.')
    p('The following traces expose longitudinal behavior, while the signal trace includes braking, a stop or release where applicable, and finite junction traversal. The physical clock is 0.1 s. Boundary remainder transfer avoids repeatedly discarding unused movement at link crossings; downstream observations maintain the leader constraint across one connector and the next road.')
    p('Those mechanisms have a specific validated scope: single-lane MICRO Cars with static routing and compatible Link clocks. This report does not extrapolate them to arbitrary multi-lane interactions, competing junction movements or very short link chains. A visually close trace is one diagnostic; terminal records, coverage and the remaining physical checks still determine acceptance.')

    page('Queues discharge and throughput')
    fig('queues-headways','Signal approach halted queue and paired discharge headways in the first measured conservative run. Queue means speed below 0.1 m/s; headways match consecutive vehicle pairs within the same green interval.')
    s=scientific.get('signal',{});q=s.get('queue',{});h=s.get('dischargeHeadways',{});tp=s.get('completedCohortThroughput',{})
    table(['Metric','Observed value'],[
        ['Queue count RMSE',number(q.get('rmseVehicles'),3)+' vehicles'],
        ['Maximum halted queue HTC and SUMO',f"{q.get('htcMaximumVehicles','NA')} and {q.get('sumoMaximumVehicles','NA')} vehicles"],
        ['Halted vehicle time HTC and SUMO',f"{number(q.get('htcVehicleSeconds'),1)} and {number(q.get('sumoVehicleSeconds'),1)} vehicle seconds"],
        ['Matched discharge pairs and headway MAE',f"{h.get('matchedPairs','NA')} pairs and {number(h.get('maeSeconds'),3)} s"],
        ['Completed cohort rate HTC and SUMO',f"{number(tp.get('htcVehiclesPerHour'),1)} and {number(tp.get('sumoVehiclesPerHour'),1)} vehicles per hour"]],[3.7,3.2])
    p('Cohort throughput divides completed vehicles by the fixed 180-second horizon. With all eighteen vehicles completing, this is a completion rate for the selected demand, not a measurement of saturated road capacity. Similar queue maxima can conceal different queue durations, so vehicle-seconds and queue curves are shown alongside the maximum.')

    page('Execution time across engines and policies')
    fig('execution-times','Individual accepted measured runs and median bars. The logarithmic axis keeps the small SUMO CLI times and HTC startup costs visible together. Warm-ups are excluded. C1 and O1 mean conservative and optimistic with one JVM; C2 and O2 use two JVMs.')
    p('The primary elapsed measurement includes process startup, cluster formation, input preparation, modeled work and the configured output pipeline. It stops before deliberate harness teardown. The historical launcher elapsedSeconds field also includes cleanup and is retained only for compatibility; it is not used for the new comparison.')
    rows=[]
    for case in CASE_ORDER[:3]:
        a=ratio(case,'discrete-event-1','sumo');b=ratio(case,'time-warp-1','sumo')
        rows.append([case.capitalize(),span(a['summary']),span(b['summary'])])
    table(['Case','Conservative 1 JVM / SUMO','Optimistic 1 JVM / SUMO'],rows,[1.4,2.75,2.75])
    p('Each ratio is calculated within a repetition block before summarization. Brackets show the observed minimum and maximum, not a confidence interval. These startup-heavy comparisons should not be read as a benchmark of a pure integration kernel or as evidence that distributed execution will pay off on a larger workload.')

    page('CPU consumption and memory footprint')
    fig('resources','CPU-seconds and peak process memory across the three traffic workloads. Dots are accepted measured runs and short bars are medians. Both axes use logarithmic scales. SUMO native RSS comes from the separately retained measurement audit; execution and CPU use the original campaign.')
    p('Two JVMs share the same total CPU allowance and maximum Java heap as one JVM, but duplicate process-level services, class metadata and thread stacks. Aggregate CPU-seconds capture that extra work even if elapsed time falls. The memory plot uses concurrent process RSS rather than allocated heap or theoretical checkpoint bounds.')
    p('CPU counters include startup, reporting, garbage collection and runtime coordination. The measurement monitor is external to these process counters, although it competes for host resources. Native SUMO RSS audit accounting and sampled Java RSS have different peak resolution, which limits fine-grained cross-engine memory comparisons.')

    page('Runtime intervals and output draining')
    fig('execution-intervals','Descriptive median intervals. Blue spans the logged simulation start request through acknowledged runtime completion. Gray is the residual launch and notification interval; it is not a pure startup measurement.')
    p('The start request can precede progressive input readiness, so the blue interval includes remaining input work as well as execution. Runtime completion already waits for committed report receipt and sink flush acknowledgements. This is a meaningful lifecycle interval, but not an isolated physics time.')
    p('The existing config_load, loading and progressive_loading timers are retained in the source records. Some phases overlap; adding their durations would overstate elapsed time. A separate output-drain duration is unavailable because the current logs do not expose both boundaries as a complete timing pair. The report therefore does not invent a standalone drain cost or subtract overlapping phases.')
    p('Comparison with SUMO uses the complete output-inclusive process path. SUMO statistics define clock duration and real-time factor for its own execution interval; that internal interval should not be substituted for the HTC launch interval in a speed ratio. [3]')

    page('Policy ratios and local scaling')
    fig('paired-ratios','Elapsed ratios paired within workload and repetition. Dots show each available pair; diamonds show medians and lines show ranges. Above one means the numerator was slower.')
    p('The optimistic-to-conservative ratios compare policies within the same JVM count. The two-to-one-JVM ratios compare deployment shape with a fixed total CPU allocation. They measure the cost or benefit of splitting these small workloads across processes on one host, not strong scaling across physical machines.')
    for case in CASE_ORDER[:3]:
        a=ratio(case,'time-warp-1','discrete-event-1')['summary'];b=ratio(case,'time-warp-2','discrete-event-2')['summary']
        p(f"{case.capitalize()}: optimistic execution takes {number(a['median'])} times the conservative elapsed time with one JVM and {number(b['median'])} times with two JVMs, paired by repetition. These are conditional accepted-run medians, with sample counts {a['n']} and {b['n']}.")
    p('Input enrollment, cross-owner transport, runtime barriers and report commitment all affect these ratios. A favorable elapsed result is useful only together with physical acceptance and report parity. A ratio based on fewer than five accepted pairs remains visible with its actual sample count in the source table.')

    page('Speculation rollback and transport')
    fig('rollbacks','Rollback log entries per accepted measured traffic run. Conservative execution has no optimistic rollback. A rollback count records speculative work that was undone, not a failed committed simulation.')
    p('Conservative execution closes the relevant causal frontier before proceeding. Optimistic execution can run ahead within fixed startup budgets, then cancel and replay work when an earlier ordered input arrives. Reports are committed only when safe, so repeated speculative work must not duplicate final physical outputs.')
    p('The source records retain each node’s transfer, receipt and cross-owner accounting summaries. Counts include optimistic cancellation and replay traffic. They prove transport activity where nonzero but are not equivalent to distinct physical interactions or useful throughput. Comparing rollback totals without elapsed and CPU costs would miss much of the execution overhead.')
    checks=summary['measuredParityChecks'];p(f'{sum(v["passed"] for v in checks)} of {len(checks)} retained measured output comparisons match their designated reference, including reference self-checks. The comparator removes only nondeterministic simulation identity, wall-clock timestamps and the summary CPU duration field; modeled physical content remains compared as a multiset.')

    page('Balancing experiment and execution cost')
    fig('balancing-times','Two-JVM conservative elapsed times with fixed spatial placement and with rebalancing enabled. Both members of each pair use the same fixture, reporting and load-observation settings.')
    rows=[]
    for case in CASE_ORDER[3:]:
        a=group(case+'-reference');b=group(case+'-migration');r=ratio(case,'migration','reference')
        rows.append([case.capitalize(),span(a['metrics']['wallSeconds']),span(b['metrics']['wallSeconds']),span(r['summary'])])
    table(['Fixture','Fixed time s','Rebalance time s','Rebalance / fixed'],rows,[1.2,1.9,1.9,1.9])
    p('The initial fixture contains sixty Persons and sixty Cars, with two trips per Person. The progressive fixture enrolls two cohorts at ticks 0 and 10050. Its 40 km roads deliberately keep Cars active across enrollment; it is a protocol stress fixture rather than a calibrated urban network. Input plans and the directed graph determine the expected endpoints, journeys and distances.')
    for case in CASE_ORDER[3:]:
        a=group(case+'-reference')['metrics'];b=group(case+'-migration')['metrics'];r=ratio(case,'migration','reference')['summary']
        if r['median'] is not None:
            change=(r['median']-1)*100
            p(f"{case.capitalize()}: the paired rebalance elapsed change has median {change:+.1f} percent. Median aggregate CPU is {number(b['cpuSeconds']['median'])} s with rebalancing versus {number(a['cpuSeconds']['median'])} s at fixed placement. The sign of the elapsed change describes this configured treatment, including its global pauses and placement effects.")
    p('Both configurations keep spatial placement and load observation enabled. The fixed reference uses a 3600-second rebalance interval, beyond the run limit; the treatment uses 0.25 seconds. These deliberately aggressive settings exercise migration and are not a recommendation for production tuning. Faster execution is not assumed.')

    page('Load distribution and imbalance over time')
    fig('live-loads','Per-node entity counts and load CV for the first measured repetition of each balancing fixture. Solid and dashed lines identify nodes; colors identify placement policy. Each run uses its own first observation as time zero.')
    p('CV equals the population standard deviation of node loads divided by their mean. For two nodes, CV is zero at equal load and one when a nonzero cohort sits entirely on one node. Maximum/mean values are retained in the raw metric record. An all-zero cohort has no defined relative imbalance.')
    p('Initialization and retirement change the observed cohort. A final one-entity remainder can have CV one even after useful balancing earlier in the run. The first and last CV therefore cannot serve as a causal estimate of balancing quality. The curves expose these changes, including periods when migration prevents a complete observation. Step lines hold the last observed value for display; values during an observation gap are not measured.')
    p('Live-entity weighting is intentionally separated from CPU. A Link performing repeated MICRO physics can cost much more than an idle entity, while these balancing fixtures use MESO roads. Per-process CPU samples and CPU-work CV are retained for examining that distinction.')

    page('CPU work imbalance during balancing')
    fig('cpu-imbalance','Per interval CPU-work CV across the two JVMs in the first measured repetition. Idle intervals are unavailable; startup and reporting remain included.')
    p('This chart compares the CPU increments of the two JVMs over successive sampled intervals. It can differ substantially from live-entity imbalance because per-entity work, coordination and reporting vary. Quarter-second intervals also expose short scheduling bursts; the plot is descriptive rather than a smoothed estimate of CPU demand.')
    rows=[]
    for case in CASE_ORDER[3:]:
        for placement in ('reference','migration'):
            g=group(case+'-'+placement)
            rows.append([case.capitalize()+' '+placement,span(g['metrics']['cpuWorkCv']),span(g['metrics']['cpuSeconds']),span(g['metrics']['peakRssMiB'])])
    table(['Configuration','Whole run CPU work CV','CPU seconds','RSS MiB'],rows,[2.2,1.7,1.5,1.5])
    p('Whole-run CPU-work CV uses the accumulated CPU counters of each process; it does not average the interval CVs. No causal CPU-balancing benefit is inferred from a single curve. The repeated elapsed and CPU comparisons provide the corresponding treatment costs.')

    page('Rebalance activity and pause costs')
    fig('rebalancing-costs','Verified physical shard moves, total completed pause-window duration and active Car restore observations. These quantities count different protocol events and must not be substituted for one another.')
    rows=[]
    for case in CASE_ORDER[3:]:
        g=group(case+'-migration');m=g['metrics']
        rows.append([case.capitalize(),span(m['verifiedShardMoves']),span(m['migrationWaves']),span(m['pauseSeconds']),span(m['activeCarRestores'])])
    table(['Fixture','Shard moves','Waves','Pause seconds','Active restores'],rows,[1.1,1.45,1.45,1.45,1.45])
    treatment=[r for r in flat if not r['warmup'] and r['kind']=='balance' and r['config'].endswith('-migration')]
    p(f"Across the ten measured rebalance treatments, the logs record {sum(r['pauseRequests'] for r in treatment)} pause requests, {sum(r['freezeDispatches'] for r in treatment)} dispatched freeze windows, {sum(r['pauseWindows'] for r in treatment)} completed resume windows and {sum(r['declinedRequests'] for r in treatment)} terminal declines. These are correlated lifecycle counts rather than counts of periodic evaluations that produced no move.")
    p('A verified shard move requires placement and snapshot-restore evidence. A migration wave can carry several shards and multiple entities. A global pause window can cover one or more waves. Active restoration additionally requires a Car whose observed road-occupancy interval contains the migration cut; an actor or shard count alone does not demonstrate that behavior.')
    p('A run that finishes without moving a shard remains evidence about execution and initial placement, but not about physical rebalancing. Counts and omissions are shown rather than replacing such runs until a move appears. Restored-entity totals count memberships per wave and can include the same entity in more than one wave.')

    page('Migration lifecycle timing')
    fig('migration-phases','Completed global pause windows in the first measured rebalance repetition. Timers use one manager’s monotonic clock and the matching protocol acknowledgements.')
    p('Freeze time ends only after all runtime owners return their frozen live cohorts. The relocation interval includes destination preparation, snapshot capture, window publication, native shard handoff, restore and placement verification, and window-close acknowledgements. Resume time ends after every owner acknowledges resumption.')
    p('These three intervals partition each completed global pause. The relocation interval is deliberately not labeled network transfer time: it includes local work and acknowledgement waiting. Individual shard-duration logs can overlap, so their sum is not a pause or wall-clock duration. Global windows execute sequentially, making their completed total meaningful within a run.')
    p('The instrumentation adds timing records at existing acknowledgement transitions. It does not create replies, reorder model events, relax completion, add watchdog recovery or make migration available under Time Warp. Failed or incomplete windows remain in raw logs and are not silently converted into completed durations.')

    page('Field evidence and empirical limits')
    image=REPO/'docs/validation/figures/micro-field-v2-2026-10-26.png'
    para=doc.add_paragraph();para.paragraph_format.keep_with_next=True
    para.add_run().add_picture(str(image),width=Inches(6.9))
    doc.add_paragraph('Retained field v2 observations and model comparisons. USDOT FHWA NGSIM Lankershim data, CC BY-SA 3.0, processed by HTC. [8] The fixed cohort targets were not met; favorable subset comparisons do not change that decision.',style='Caption')
    rows=[]
    for kind in ('following','queue'):
        s=field['primarySummaries'][kind]
        rows.append([kind.capitalize(),f"{field['selectionCounts'][kind]} / {field['targets'][kind]}",number(s['positionRmseM']['mean'],3),number((s.get('headwayMaeSeconds') or {}).get('mean'),3)])
    table(['Cohort','Qualified / required','Position RMSE m','Headway MAE s'],rows,[1.5,1.8,1.8,1.8])
    p('The field replay drives the observed head leader and predicts followers using the real Link physics. Its same-section following scope, observed release alignment and posted-speed input make it a conditional longitudinal assessment. It does not reconstruct a measured green phase or validate an autonomous network. The supporting signal attachment failed integrity checks.')
    p('All comparative checks pass on the qualified v2 subset, with complete predicted crossings and no overlap. The frozen sample-size gate still yields not-established. Speed-cap sensitivities produce identical paths in this leader-dominated subset and therefore do not identify individual desired speeds. No coefficients are retuned for this report.')

    page('Historical validation and retained failures')
    table(['Evidence','Supported result','Boundary'],[
        ['SUMO v1 scientific study','Physical cases passed; five randomized-demand cases failed their frozen criteria.','Historical failures remain development evidence.'],
        ['SUMO v2 cross-link following','Four physical cases and five fresh demand seeds passed after the boundary-observation correction.','Matched model inputs; not field validity.'],
        ['Multi-JVM policies','Committed-output parity and physical outcomes established on bounded road and transit cases.','Loopback processes; historical raw paths partly unavailable.'],
        ['Conservative migration','Verified physical restore and output parity on initial and progressive cohorts.','Optimistic migration and independent recovery unsupported.'],
        ['MICRO spillback','Bounded single-lane admission and retained-footprint invariants passed.','General junction competition and network spillback remain open.'],
        ['Field v2','Eight following pairs and four queues pass subset comparisons.','Targets ten and five unmet; empirical gate inconclusive.']], [1.5,3,2.4])
    p('The v1 dense-demand failures diagnosed a missing leader constraint across link boundaries. The later v2 cases used a frozen new protocol and fresh geometries/seeds; v1 was not retrospectively relabeled as passing. A v2 optimistic stale-publication engine failure was repaired without changing the traffic criteria, then the complete campaign was repeated. That initial failure remains excluded engine evidence.')
    p('Historical summaries retain executable and protocol identities. Some temporary raw directories cited in older documents are no longer present on this host. Their tables are attributed to retained machine records, rather than represented as freshly revalidated raw observations. New campaign evidence is archived with hashes and actual commands. The initial input-preparation attempt failed because multiple generator calls shared an entity-ID registry; resetting that registry between independent fixtures corrected generation before the manifest was frozen. This was a preparation failure, not a traffic result or a reason to alter physical thresholds.')

    page('Limitations and conclusions for further work')
    p('This report separates three questions: whether HTC preserves modeled outcomes across runtime policies; whether selected HTC traffic traces agree with matched SUMO cases; and what the configured executions cost on this host. Evidence for one does not establish the others. Empirical validity remains open despite favorable subset and SUMO comparisons.')
    p('Small cohorts make startup, cluster formation and detailed reporting a large fraction of the measured path. The benchmark does not explore saturated demand, large actor counts, multi-machine networks, membership growth, dynamic routing overhead or failure recovery. Five deterministic repetitions describe local variability, not a general performance distribution.')
    p('Custom rebalancing remains disabled by default. The treatment globally pauses conservative execution, migrates state and waits for all owner and placement acknowledgements. A more balanced entity distribution can still cost more CPU or elapsed time. The report supports using the measured tradeoff for these fixtures, while leaving production tuning and city-scale benefits unestablished.')
    p('Optimistic execution can be selected for supported workloads without migration; its speculative budgets are fixed at startup and do not impose a hard heap-byte cap. Output parity and rollback behavior are necessary checks before interpreting runtime gains. Optimistic migration is outside the supported matrix.')
    p('The next empirical study needs a newly preregistered cohort of uninspected observations with enough sampling opportunities to meet the original targets. Acceptance limits should not be lowered and favorable inspected cases should not be pooled after their errors are known. Performance findings in this report do not justify changing the traffic coefficients.')
    heading('Completion and exceptions')
    p(f'All {summary["attempts"]} planned attempts are accounted for. {summary["acceptedMeasured"]} of {summary["measuredAttempts"]} measured attempts are accepted for their designated checks. The 18 additional RSS-only audit attempts are accounted for separately and reproduce the original SUMO trip records exactly. Missing lifecycle timings and unsupported configurations remain explicit. The appendices retain per-attempt values so readers can examine exclusions and variability directly.')
    if summary['failedMeasured']:
        for r in summary['failedMeasured']:p(f"Measured exception r{r['repeat']} {r['config']}: {r.get('failure') or 'See runtime, physical and parity checks in the evidence ledger.'}")

    page('Reproduction and evidence index')
    p('The accompanying machine tables and evidence archive are the source for every newly plotted value. Configuration IDs map directly to case, policy, JVM count or placement. r0 denotes excluded warm-up, and r1 through r5 are measured repetitions. The first measured repetition supplies the illustrative time series, while distribution plots use all accepted measured runs.')
    table(['Artifact','Purpose'],[
        ['manifest.json and manifest-seal.json','Frozen plan, source/input hashes, executable identity, software and host metadata.'],
        ['results.json','Every planned attempt, acceptance checks and report-parity outcomes.'],
        ['all-attempts.csv and JSON','Plotting-ready measured values, warm-up flags, failures and raw source paths.'],
        ['configuration-summary.csv and summary.json','Accepted sample counts, medians, ranges and paired ratios.'],
        ['Raw evidence archive','Inputs, commands, logs, reports, process/Prometheus samples and checksums.'],
        ['Figures PNG and SVG','Standalone charts with the same values used in the report.']], [2.6,4.3])
    p('Reproduction uses scripts/scenario/run_comparison_study.py with --freeze followed by --run in a fresh writable directory. Build the current executable with sbt test assembly first. Use the retained manifest to recover the exact inputs and implementation identity; an arbitrary later build is a new comparison. Run --finalize to rebuild acceptance and parity summaries without executing simulations.')
    p('The report source is scripts/scenario/build_comparison_report.py. Its --analyze stage derives tables and scientific figures from the final result ledger; --docx writes the editable report with the bundled document runtime. Rendering produces the corresponding PDF and page images for visual inspection. Public output documents contain no fitted trace alignment or silently removed failures.')
    p('Manifest SHA256: '+summary['manifestSha256'])
    p('Executable SHA256: '+manifest['jarSha256'])
    heading('References')
    for text in [
        '[1] Eclipse SUMO. TripInfo documentation. https://sumo.dlr.de/docs/Simulation/Output/TripInfo.html. Accessed 27 October 2026.',
        '[2] Eclipse SUMO. FCDOutput documentation. https://sumo.dlr.de/docs/Simulation/Output/FCDOutput.html. Accessed 27 October 2026.',
        '[3] Eclipse SUMO. Statistic Output documentation. https://sumo.dlr.de/docs/Simulation/Output/StatisticOutput.html. Accessed 27 October 2026.',
        '[4] HTC repository. MICRO_CROSS_LINK_FOLLOWING.md and validation/micro-discharge-2026-10-25.json. Frozen v2 numerical comparisons.',
        '[5] HTC repository. DISTRIBUTED_LOAD_BALANCING_VALIDATION_2026_10_26.md and validation/load-balancing-completion-2026-10-26.json.',
        '[6] HTC repository. MICRO_FIELD_INPUT_AUDIT.md and validation/micro-field-v2-2026-10-26.json. Conditional field assessment.',
        '[7] HTC repository. MULTI_JVM_POLICY_VALIDATION_2026_10_24.md and MICRO_SPILLBACK.md. Historical runtime and boundary acceptance.',
        '[8] USDOT FHWA. Next Generation Simulation NGSIM vehicle trajectories, Lankershim Boulevard. https://data.transportation.gov/Automobiles/Next-Generation-Simulation-NGSIM-Vehicle-Trajectories/8ect-6jqj. CC BY-SA 3.0; retained observation snapshots contain source queries and licensing.'
    ]:
        para=p(text);para.paragraph_format.space_after=Pt(4);para.paragraph_format.line_spacing=1
        for run in para.runs:run.font.size=Pt(9.5)

    page('Measured configuration summary')
    p('Accepted measured runs only. Values are medians followed by the observed minimum and maximum. All configurations have five scheduled measured attempts; the accepted count reveals any exclusion. CPU is aggregate process CPU-seconds and memory is aggregate peak RSS in MiB.')
    rows=[]
    for case in CASE_ORDER:
        for g in summary['groups']:
            if g['case']!=case:continue
            key=g['id'].split('-',1)[1];m=g['metrics']
            label={'sumo':'SUMO','discrete-event-1':'C1','time-warp-1':'O1','discrete-event-2':'C2','time-warp-2':'O2','reference':'Fixed','migration':'Rebalance'}[key]
            rows.append([case.capitalize()+' '+label,f"{g['accepted']}/{g['attempts']}",span(m['wallSeconds']),span(m['cpuSeconds']),span(m['peakRssMiB'])])
    table(['Configuration','Accepted','Elapsed s','CPU s','RSS MiB'],rows,[1.8,.7,1.47,1.47,1.46])

    page('Individual measured attempts')
    p('Every measured attempt is listed below. Warm-ups remain in the machine ledger. A failed acceptance appears explicitly and does not enter the distribution summaries. Time is launch-to-completion before cleanup; an unavailable value is never replaced with zero.')
    rows=[]
    for r in sorted((v for v in flat if not v['warmup']),key=lambda v:(CASE_ORDER.index(v['case']),v['config'],v['repeat'])):
        rows.append([r['config'],r['repeat'],'Pass' if r['accepted'] else 'Fail',number(r['wallSeconds']),number(r['cpuSeconds']),number(r['peakRssMiB'],1),r['verifiedShardMoves']])
    table(['Configuration','Rep','Check','Time s','CPU s','RSS MiB','Moves'],rows,[2.4,.4,.55,.85,.85,.95,.6])
    for style in doc.styles:
        if style.type == 1:
            pr=style.element.find(qn('w:pPr'))
            if pr is not None:
                for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
        if style.name in ('Title','Subtitle','Heading 1','Heading 2','Heading 3','Normal','Caption'):
            rpr=style.element.get_or_add_rPr()
            fonts=rpr.find(qn('w:rFonts'))
            if fonts is None:fonts=OxmlElement('w:rFonts');rpr.append(fonts)
            for attribute in list(fonts.attrib):
                if 'Theme' in attribute:del fonts.attrib[attribute]
            for key in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+key),'Times New Roman')
    for para in doc.paragraphs:
        pr=para._p.find(qn('w:pPr'))
        if pr is not None:
            for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
    doc.core_properties.title='HTC traffic fidelity and runtime performance'
    doc.core_properties.subject='Matched SUMO traffic comparisons, conservative and optimistic execution, conservative rebalancing'
    doc.core_properties.author='HTC Digital Twin Platform'
    doc.core_properties.keywords='HTC SUMO Time Warp conservative optimistic load balancing benchmark'
    path=out/'htc-comparison-report-2026-10-27.docx';doc.save(path)
    print(path)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--study',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--analyze',action='store_true')
    parser.add_argument('--docx',action='store_true')
    args=parser.parse_args();study=args.study.resolve();out=args.output.resolve()
    if args.analyze:
        summary,flat,manifest,data=derive(study,out)
        saved=figures(study,out,summary,flat,manifest,data)
        write(out/'figure-index.json',{name:{ext:digest(out/'figures'/(name+'.'+ext)) for ext in ('png','svg')} for name in saved})
    if args.docx:document(study,out)
